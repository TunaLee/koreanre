#!/usr/bin/env python3
"""오전 실습용 더미 자료를 data/오전/ 에 만든다. 모든 이름 · 숫자는 가상.

사용: python3 src/make_data.py --week 2026-10-26   (교육 주의 월요일)
캘린더 일정과 마감일이 그 주를 기준으로 잡히므로, 교육 날짜가 정해지면 다시 만든다.
"""
import argparse
import random
from datetime import date, datetime, timedelta
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "오전"

TASKS = [  # 업무명, 담당, 마감(주 시작일 기준 일수), 진행률, 비고
    ("사내 교육 일정 확정", "김서연", -2, 90, "강사 1명 회신 대기"),
    ("3분기 비용 집계", "이준호", 1, 100, ""),
    ("자료 보관 기준 정비", "박지민", 3, None, "담당 부서 협의 중"),
    ("협력사 명단 갱신", "최유진", 4, 60, ""),
    ("연말 행사 장소 섭외", "정하늘", -5, 40, "후보 3곳 중 2곳 회신"),
    ("부서 비품 구매 요청", "한도윤", 2, 100, ""),
    ("업무 매뉴얼 개정", "김서연", 9, 30, "목차 확정"),
    ("신규 입사자 온보딩 자료", "윤채원", 4, 70, ""),
    ("월간 뉴스레터 발행", "이준호", 3, 80, "디자인 검토 중"),
    ("보안 점검 체크리스트 회수", "박지민", 1, 50, "12개 팀 중 6개 팀 회수"),
    ("출장비 정산 취합", "최유진", 8, 20, ""),
    ("회의실 예약 규칙 안내", "한도윤", 0, 100, "공지 완료"),
]

EVENTS = [  # 요일(0=월), 시작, 끝, 제목, 장소
    (0, "09:30", "10:00", "주간 팀 회의", "7층 회의실"),
    (0, "14:00", "15:00", "교육 일정 강사 통화", "온라인"),
    (1, "10:00", "11:30", "협력사 미팅", "외근 · 광화문"),
    (1, "16:00", "16:30", "비용 집계 결과 보고", "팀장실"),
    (2, "13:30", "14:30", "자료 보관 기준 부서 협의", "5층 회의실"),
    (2, "17:00", "17:30", "뉴스레터 디자인 검토", "온라인"),
    (3, "10:00", "12:00", "신규 입사자 오리엔테이션", "대회의실"),
    (3, "15:00", "16:00", "행사 장소 답사", "외근 · 종로"),
    (4, "11:00", "11:30", "보안 점검 회수 마감 확인", "7층 회의실"),
    (4, "16:00", "17:00", "주간 업무 정리", "자리"),
]


def head(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="EEEEEE")


def sheet_tasks(week):
    wb = Workbook()
    ws = wb.active
    ws.title = "팀 업무 현황"
    head(ws, ["업무명", "담당", "마감일", "진행률(%)", "비고"])
    for name, who, d, pct, memo in TASKS:
        ws.append([name, who, (week + timedelta(days=d)).isoformat(), pct, memo])
    for col, w in zip("ABCDE", (26, 10, 12, 11, 30)):
        ws.column_dimensions[col].width = w
    wb.save(OUT / "팀_업무_현황.xlsx")


def ics(week):
    lines = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//training//KO", "CALSCALE:GREGORIAN", "X-WR-CALNAME:실습용 팀 일정"]
    stamp = datetime.now().strftime("%Y%m%dT%H%M%SZ")
    for i, (wd, s, e, title, loc) in enumerate(EVENTS):
        day = week + timedelta(days=wd)
        f = lambda t: day.strftime("%Y%m%d") + "T" + t.replace(":", "") + "00"
        lines += ["BEGIN:VEVENT", f"UID:team-{i}-{day.isoformat()}@training", f"DTSTAMP:{stamp}",
                  f"DTSTART;TZID=Asia/Seoul:{f(s)}", f"DTEND;TZID=Asia/Seoul:{f(e)}",
                  f"SUMMARY:{title}", f"LOCATION:{loc}", "END:VEVENT"]
    lines.append("END:VCALENDAR")
    (OUT / "실습용_팀일정.ics").write_text("\r\n".join(lines) + "\r\n", encoding="utf-8")


def docs(week):
    d = Document()
    d.add_heading("주간 업무 보고", 0)
    d.add_paragraph("작성 주: YYYY년 MM월 N주차 · 작성자: ○○○")
    d.add_heading("1. 이번 주 요약", 1)
    d.add_paragraph("(한 문장)")
    d.add_heading("2. 업무별 현황", 1)
    t = d.add_table(rows=2, cols=5)
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ["업무명", "담당", "진행률", "다음 조치", "비고"]):
        c.text = h
    d.add_heading("3. 이번 주 주요 일정", 1)
    d.add_paragraph("- MM/DD (요일) 시간 · 일정명")
    d.save(OUT / "주간보고_양식.docx")

    prev = week - timedelta(days=7)
    d = Document()
    d.add_heading("주간 업무 보고", 0)
    d.add_paragraph(f"작성 주: {prev.year}년 {prev.month}월 {(prev.day - 1) // 7 + 1}주차 · 작성자: 김서연")
    d.add_heading("1. 이번 주 요약", 1)
    d.add_paragraph("11건 중 완료 3건 · 마감 초과 1건(사내 교육 일정 확정, 강사 회신 대기) · 확인 요청 1건.")
    d.add_heading("2. 업무별 현황", 1)
    rows = [("[마감 2일 초과] 사내 교육 일정 확정", "김서연", "80%", "미회신 강사 1명 재연락", "강사 회신 대기"),
            ("3분기 비용 집계", "이준호", "70%", "부서별 증빙 취합", ""),
            ("자료 보관 기준 정비", "박지민", "확인 요청", "진행률 확인 필요", "담당 부서 협의 중"),
            ("협력사 명단 갱신", "최유진", "40%", "변경 업체 연락처 확인", "")]
    t = d.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ["업무명", "담당", "진행률", "다음 조치", "비고"]):
        c.text = h
    for r in rows:
        for c, v in zip(t.add_row().cells, r):
            c.text = v
    d.add_heading("3. 이번 주 주요 일정", 1)
    d.add_paragraph("- 월 09:30 주간 팀 회의\n- 수 13:30 비용 집계 중간 점검\n- 금 16:00 주간 업무 정리")
    d.save(OUT / "작성예시.docx")

    for i, (subj, body) in enumerate([
        ("[주간 보고] 이번 주 팀 업무 현황 공유드립니다",
         "팀장님, 이번 주 팀 업무 현황 공유드립니다.\n\n요약: 11건 중 완료 3건, 마감 초과 1건입니다.\n"
         "마감 초과 건은 강사 회신 대기 중이며 수요일까지 재연락하겠습니다.\n\n상세 내용은 첨부한 주간 보고서를 확인 부탁드립니다.\n\n김서연 드림"),
        ("[확인 요청] 자료 보관 기준 정비 진행률 회신 부탁드립니다",
         "박지민 님, 주간 보고 작성 중 「자료 보관 기준 정비」 진행률이 비어 있어 연락드립니다.\n"
         "목요일 오전까지 현재 진행률과 다음 조치를 회신해 주시면 반영하겠습니다.\n\n감사합니다.\n김서연 드림"),
    ], 1):
        d = Document()
        d.add_paragraph(f"제목: {subj}")
        for para in body.split("\n"):
            d.add_paragraph(para)
        d.save(OUT / f"메일예시_{i}.docx")


def examples(week):
    ex = OUT / "예시3종"
    ex.mkdir(exist_ok=True)
    r = random.Random(3)
    wb = Workbook()
    ws = wb.active
    ws.title = "경비 내역"
    head(ws, ["사용일", "항목", "내용", "금액(원)", "사용자"])
    items = {"교통비": ["택시", "KTX", "주차"], "식대": ["팀 점심", "야근 식대", "손님 접대"], "비품": ["토너", "문구", "모니터 암"],
             "교육비": ["외부 세미나", "온라인 강의"], "회의비": ["다과", "회의실 대관"]}
    names = ["김서연", "이준호", "박지민", "최유진", "정하늘", "한도윤"]
    for _ in range(30):
        cat = r.choice(list(items))
        ws.append([(week - timedelta(days=r.randint(1, 28))).isoformat(), cat, r.choice(items[cat]),
                   r.choice([8000, 12500, 23000, 45000, 59800, 120000, 18000, 34000, 6500]), r.choice(names)])
    b = wb.create_sheet("월 예산")
    head(b, ["항목", "월 예산(원)"])
    for cat, v in [("교통비", 200000), ("식대", 300000), ("비품", 150000), ("교육비", 200000), ("회의비", 100000)]:
        b.append([cat, v])
    wb.save(ex / "경비_내역.xlsx")

    d = Document()
    d.add_heading("회의 메모 (정리 전)", 1)
    for line in ["주간 팀 회의 / 7층 회의실 / 참석: 김서연 이준호 박지민 최유진 정하늘",
                 "- 연말 행사 장소 후보 3곳 중 1곳 탈락, 나머지 2곳 답사 목요일. 정하늘 담당",
                 "- 뉴스레터 이번 달은 디자인 바꾸기로 함. 이준호가 시안 2개 수요일까지",
                 "- 보안 점검 체크리스트 절반만 들어옴. 박지민이 미회수 팀에 다시 연락. 금요일 마감",
                 "- 비품 구매는 승인 완료, 다음 주 도착 예정",
                 "- 다음 회의 때 매뉴얼 개정 목차 공유 (김서연)",
                 "- 출장비 정산 양식 바뀐 거 공지해야 함 → 누가? 미정"]:
        d.add_paragraph(line)
    d.save(ex / "회의_메모.docx")

    wb = Workbook()
    ws = wb.active
    ws.title = "프로젝트 일정"
    head(ws, ["작업", "담당", "시작일", "마감일", "상태"])
    for name, who, s, e, st in [("요구사항 정리", "김서연", -20, -10, "완료"), ("화면 설계", "박지민", -12, -2, "진행 중"),
                                ("업체 견적 비교", "최유진", -7, 2, "진행 중"), ("내부 결재", "이준호", 1, 4, "대기"),
                                ("계약 체결", "이준호", 5, 9, "대기"), ("사용자 안내 자료", "윤채원", 0, 4, "진행 중"),
                                ("시범 운영", "한도윤", 10, 20, "대기"), ("교육 일정 공지", "정하늘", -3, 0, "진행 중")]:
        ws.append([name, who, (week + timedelta(days=s)).isoformat(), (week + timedelta(days=e)).isoformat(), st])
    wb.save(ex / "프로젝트_일정.xlsx")


def main():
    ap = argparse.ArgumentParser()
    nxt = date.today() + timedelta(days=(7 - date.today().weekday()) % 7 or 7)
    ap.add_argument("--week", default=nxt.isoformat(), help="교육 주 월요일 (YYYY-MM-DD)")
    week = date.fromisoformat(ap.parse_args().week)
    OUT.mkdir(parents=True, exist_ok=True)
    sheet_tasks(week)
    ics(week)
    docs(week)
    examples(week)
    print("기준 주:", week, "→", sorted(p.name for p in OUT.rglob("*") if p.is_file()))


if __name__ == "__main__":
    main()
